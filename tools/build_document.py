from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
root=Path(__file__).resolve().parents[1]
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.72)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Header','Footer','Caption']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(10.5 if name=='Normal' else 28 if name=='Title' else 19 if name=='Heading 1' else 13 if name=='Heading 2' else 10)
 st.paragraph_format.space_after=Pt(7)
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(13);st.paragraph_format.keep_with_next=True
normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=1.08
for name in ['List Bullet','List Number']:
 doc.styles[name].font.name='Arial';doc.styles[name].font.size=Pt(10.5)
 doc.styles[name].paragraph_format.space_after=Pt(5)
header=sec.header.paragraphs[0];header.text='BOX COUNTING  /  BROWSER GAME PLAN';header.style='Header';header.runs[0].font.size=Pt(8)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('13 October 2026  •  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

def inline(p,text):
 for part in re.split(r'(https?://\S+|\*\*.*?\*\*)',text):
  if part.startswith('http'):
   rel=p.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
   h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'24546B');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=part;r.append(t);h.append(r);p._p.append(h)
  else:
   run=p.add_run(part[2:-2] if part.startswith('**') else part)
   run.bold=part.startswith('**')

def table(rows):
 n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if n==7: ratios=[.08,.115,.11,.11,.09,.15,.345]
 elif n==6:ratios=[.1,.12,.12,.16,.16,.34]
 elif n==4:ratios=[.19,.32,.15,.34]
 elif n==3:ratios=[.075,.435,.49] if rows[0][0]=='ID' else [.24,.38,.38]
 else:ratios=[1/n]*n
 for col,r in zip(t.columns,ratios):col.width=Inches(7.06*r)
 for i,row in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,text in enumerate(row):
   c=cells[j];c.width=Inches(7.06*ratios[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8EE' if i==0 else 'FFFFFF');pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders);m=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');m.append(el)
   pr.append(m);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
   run=p.add_run(text);run.font.size=Pt(8.8 if n>=6 else 9.2);run.bold=i==0
  trpr=t.rows[i]._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
  if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)

def figure(path,caption,width=6.2):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 r=p.add_run();pic=r.add_picture(str(root/path),width=Inches(width));pic._inline.docPr.set('descr',caption)
 p=doc.add_paragraph(caption,style='Caption');p.paragraph_format.space_after=Pt(12)

files=sorted((root/'docs').glob('0[0-6]-*.md'))
for idx,path in enumerate(files):
 # Continue chapters naturally to avoid nearly empty chapter-ending pages.
 lines=path.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.match(r'^:?-+:?$',x) for x in cells):rows.append(cells)
    i+=1
   table(rows);continue
  if line.startswith('# '):
   doc.add_paragraph(line[2:],'Title' if idx==0 else 'Heading 1')
  elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 2')
  elif re.match(r'^\d+\. ',line):inline(doc.add_paragraph(),line)
  elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
  else:inline(doc.add_paragraph(),line)
  i+=1
# Append genuine frames and proposed visuals with explicit provenance.
doc.add_page_break();doc.add_heading('Reference evidence figures',1)
doc.add_paragraph('These are frames extracted from V3. They establish details that the original pasted specification missed.')
figure('references/frames/large-boxes.png','V3 at about 5 minutes 52 seconds. The broad box is one object; the full arrangement is four boxes.',5.8)
figure('references/frames/final-score.png','V3 at 6 minutes 45 seconds. Time plus five seconds per mistake determines the winner.',5.8)
doc.add_page_break();doc.add_heading('Browser visual targets',1)
doc.add_paragraph('Newly drawn geometric targets for the browser version. Exact positions, palette and font remain subject to implementation comparison.')
figure('assets/visual-targets/watch.png','Watch target. Neutral boxes, faint inactive counters and a fixed isometric projection.',5.8)
figure('assets/visual-targets/answer.png','Answer target. The boxes are gone and both tiled counters accept independent keyboard input.',5.8)
doc.add_page_break();doc.add_heading('Reveal visual target',1)
figure('assets/visual-targets/reveal.png','Final reveal target. All six boxes are green; the submitted six receives a check and the seven a cross.',6.8)
doc.add_paragraph('During the implemented reveal animation, boxes turn green one object at a time and the central number rises in the same order. This still image represents the final state. The cyan and coral badges replace Switch controller illustrations.')
doc.add_heading('Supporting files',2)
doc.add_paragraph('The Markdown chapters, timestamped evidence, downloaded video references, editable SVGs, rules profile, schema and ten-level reconstruction pack live beside this document in the project folder. See README.md for the complete index. The playable release, browser checks, and subsequent video/mode expansion are documented in docs/08-playable-prototype.md, docs/09-release-qa.md, and docs/10-video-layouts-and-endless-duel.md. This Word export is a snapshot of the planning chapters.')
doc.core_properties.title='Box Counting browser recreation plan';doc.core_properties.subject='Gameplay visual fidelity levels and implementation';doc.core_properties.author='';doc.core_properties.keywords='Box Counting, local multiplayer, browser game, implementation plan'
for el in list(doc.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in list(doc.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
out=root/'output/Box-Counting-Game-Plan.docx';doc.save(out)
print(out)
